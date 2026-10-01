"""Tests for the mihcsme CLI."""

import openpyxl
import pytest

pytest.importorskip("typer")

from typer.testing import CliRunner  # noqa: E402

from mihcsme_py.cli import app  # noqa: E402
from mihcsme_py.models import (  # noqa: E402
    DataOwner,
    InvestigationInformation,
    MIHCSMEMetadata,
)
from mihcsme_py.writer import DEFAULT_TEMPLATE_PATH  # noqa: E402

runner = CliRunner()


@pytest.fixture
def json_file(tmp_path):
    metadata = MIHCSMEMetadata(
        investigation_information=InvestigationInformation(
            data_owner=DataOwner(first_name="Jane", last_name="Doe")
        )
    )
    path = tmp_path / "metadata.json"
    path.write_text(metadata.model_dump_json(by_alias=True))
    return path


def _first_name(path):
    ws = openpyxl.load_workbook(path)["InvestigationInformation"]
    return next(
        row[2] for row in ws.iter_rows(values_only=True) if row[1] == "First Name"
    )


def test_to_excel_plain(json_file, tmp_path):
    out = tmp_path / "out.xlsx"
    result = runner.invoke(app, ["to-excel", str(json_file), "-o", str(out)])
    assert result.exit_code == 0, result.output
    wb = openpyxl.load_workbook(out)
    assert not any(name.startswith("_") for name in wb.sheetnames)
    assert _first_name(out) == "Jane"


def test_to_excel_template(json_file, tmp_path):
    out = tmp_path / "out.xlsx"
    result = runner.invoke(app, ["to-excel", str(json_file), "--template", "-o", str(out)])
    assert result.exit_code == 0, result.output
    assert openpyxl.load_workbook(out).sheetnames == (
        openpyxl.load_workbook(DEFAULT_TEMPLATE_PATH).sheetnames
    )
    assert _first_name(out) == "Jane"


def test_to_excel_template_file_implies_template(json_file, tmp_path):
    out = tmp_path / "out.xlsx"
    result = runner.invoke(
        app,
        ["to-excel", str(json_file), "--template-file", str(DEFAULT_TEMPLATE_PATH), "-o", str(out)],
    )
    assert result.exit_code == 0, result.output
    assert "_efo_assaytypes" in openpyxl.load_workbook(out).sheetnames
    assert _first_name(out) == "Jane"
